"""Document normalization layer.

Transforms multi-format inputs (PDF, JPEG, DOCX, Text) into a standard
internal representation ready for Gemini multimodal input or text analysis.
"""

from __future__ import annotations

import base64
import io
from dataclasses import dataclass
from typing import Optional
import docx
from core.exceptions import InputValidationError
from input.validator import validate_file_payload, validate_pasted_text


@dataclass
class NormalizedDocument:
    """Normalized payload representation for Gemini classification and extraction."""
    filename: str
    mime_type: str
    is_multimodal: bool
    raw_text: Optional[str] = None
    base64_data: Optional[str] = None
    file_size_bytes: int = 0

    def to_dict(self) -> dict:
        """Convert to dictionary for ephemeral session caching."""
        return {
            "filename": self.filename,
            "mime_type": self.mime_type,
            "is_multimodal": self.is_multimodal,
            "raw_text": self.raw_text,
            "base64_data": self.base64_data,
            "file_size_bytes": self.file_size_bytes,
        }

    @classmethod
    def from_dict(cls, data: dict) -> NormalizedDocument:
        """Reconstruct NormalizedDocument from cached dictionary."""
        return cls(
            filename=data["filename"],
            mime_type=data["mime_type"],
            is_multimodal=data["is_multimodal"],
            raw_text=data.get("raw_text"),
            base64_data=data.get("base64_data"),
            file_size_bytes=data.get("file_size_bytes", 0),
        )


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract plain text from DOCX file bytes including paragraphs and tables."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = []

        # Extract headings and paragraphs
        for para in doc.paragraphs:
            stripped = para.text.strip()
            if stripped:
                text_parts.append(stripped)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_parts.append(" | ".join(row_cells))

        full_text = "\n\n".join(text_parts).strip()
        if not full_text:
            raise InputValidationError("The DOCX document appears to be empty or contains no readable text.")

        return full_text
    except Exception as exc:
        if isinstance(exc, InputValidationError):
            raise
        raise InputValidationError(f"Unable to parse DOCX content: {str(exc)}") from exc


def normalize_uploaded_file(file_bytes: bytes, filename: str) -> NormalizedDocument:
    """Validate and normalize an uploaded file into a standard NormalizedDocument."""
    mime_type, ext = validate_file_payload(file_bytes, filename)
    size_bytes = len(file_bytes)

    if ext == ".docx":
        extracted_text = extract_text_from_docx(file_bytes)
        return NormalizedDocument(
            filename=filename,
            mime_type="text/plain",
            is_multimodal=False,
            raw_text=extracted_text,
            file_size_bytes=size_bytes,
        )

    elif ext == ".txt":
        text_content = file_bytes.decode("utf-8").strip()
        return NormalizedDocument(
            filename=filename,
            mime_type="text/plain",
            is_multimodal=False,
            raw_text=text_content,
            file_size_bytes=size_bytes,
        )

    elif ext in (".pdf", ".jpg", ".jpeg"):
        encoded_b64 = base64.b64encode(file_bytes).decode("utf-8")
        return NormalizedDocument(
            filename=filename,
            mime_type=mime_type,
            is_multimodal=True,
            base64_data=encoded_b64,
            file_size_bytes=size_bytes,
        )

    raise InputValidationError(f"Unsupported document extension: {ext}")


def normalize_pasted_text(text: str, filename: str = "pasted_contract.txt") -> NormalizedDocument:
    """Validate and normalize direct pasted text."""
    clean_text = validate_pasted_text(text)
    return NormalizedDocument(
        filename=filename,
        mime_type="text/plain",
        is_multimodal=False,
        raw_text=clean_text,
        file_size_bytes=len(clean_text.encode("utf-8")),
    )
