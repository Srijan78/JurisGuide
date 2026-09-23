"""Unit tests for input layer validation, security checks, and normalization."""

import io
import pytest
import docx
from core.exceptions import InputValidationError, FileSizeLimitExceededError
from input.validator import validate_file_payload, validate_pasted_text
from input.normalizer import normalize_uploaded_file, normalize_pasted_text, NormalizedDocument


def test_valid_pdf_magic_bytes():
    pdf_bytes = b"%PDF-1.7\n%stream\ncontent\n%%EOF"
    mime, ext = validate_file_payload(pdf_bytes, "document.pdf")
    assert mime == "application/pdf"
    assert ext == ".pdf"


def test_spoofed_pdf_rejected():
    fake_pdf = b"MZ\x90\x00BinaryContent"
    with pytest.raises(InputValidationError) as exc:
        validate_file_payload(fake_pdf, "evil.pdf")
    assert "%PDF" in exc.value.message


def test_valid_jpeg_magic_bytes():
    jpg_bytes = b"\xff\xd8\xff\xe0\x00\x10JFIF"
    mime, ext = validate_file_payload(jpg_bytes, "scan.jpg")
    assert mime == "image/jpeg"
    assert ext == ".jpg"


def test_spoofed_jpeg_rejected():
    fake_jpg = b"NotAJPEGHeader"
    with pytest.raises(InputValidationError) as exc:
        validate_file_payload(fake_jpg, "fake.jpeg")
    assert "JPEG specification" in exc.value.message


def test_valid_docx_extraction():
    stream = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph("Legal Employment Offer Letter")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Notice Period"
    table.rows[0].cells[1].text = "30 Days"
    doc.save(stream)

    normalized = normalize_uploaded_file(stream.getvalue(), "offer.docx")
    assert normalized.is_multimodal is False
    assert "Legal Employment Offer Letter" in normalized.raw_text
    assert "Notice Period | 30 Days" in normalized.raw_text


def test_corrupted_docx_rejected():
    fake_zip = io.BytesIO()
    import zipfile
    with zipfile.ZipFile(fake_zip, "w") as zf:
        zf.writestr("test.txt", "not a docx")
    with pytest.raises(InputValidationError) as exc:
        normalize_uploaded_file(fake_zip.getvalue(), "fake.docx")
    assert "Missing internal document structure" in exc.value.message


def test_empty_file_rejected():
    with pytest.raises(InputValidationError) as exc:
        validate_file_payload(b"", "empty.pdf")
    assert "empty" in exc.value.message


def test_file_size_limit_exceeded():
    large_payload = b"%PDF" + (b"0" * (5 * 1024 * 1024 + 1))
    with pytest.raises(FileSizeLimitExceededError):
        validate_file_payload(large_payload, "large.pdf")


def test_unsupported_file_extension():
    with pytest.raises(InputValidationError) as exc:
        validate_file_payload(b"some content", "script.sh")
    assert "Unsupported file format" in exc.value.message


def test_pasted_text_valid_and_sanitized():
    text = "  This is an agreement of service between party A and party B with sufficient length.  "
    norm = normalize_pasted_text(text)
    assert norm.raw_text == text.strip()
    assert norm.is_multimodal is False


def test_pasted_text_too_short_rejected():
    with pytest.raises(InputValidationError) as exc:
        validate_pasted_text("Too short text")
    assert "too brief" in exc.value.message


def test_pasted_text_empty_rejected():
    with pytest.raises(InputValidationError):
        validate_pasted_text("   \n\t  ")
